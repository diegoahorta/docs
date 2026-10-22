#!/usr/bin/env python3
"""Extrai, valida e mapeia os vocabulários ODU Creative / Mnemotechnics (.docx).

Uso: python3 -I extrair.py <pasta-com-docx> <pasta-de-saida>

Fonte mestre: Vocabulario_Essencial_Multilingue_2284 (EN + 11 idiomas).
Os demais arquivos são cruzados com ela e anexados como camadas extras:
  - volumes 01/02/04/06 (EN-PT, EN-ES, EN-FR, EN-DE): alertas + sugestões
  - Alta_Frequencia_Ingles_2000/Expandido/Pedagogico: tradução PT-BR curada
    e frases de exemplo.
"""
import csv, json, re, sys
from pathlib import Path

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

LANGS = ["PT", "ES", "IT", "FR", "RO", "DE", "NL", "SV", "NO", "DA", "IS"]
LANG_NAMES = {
    "EN": "Inglês", "PT": "Português", "ES": "Espanhol", "IT": "Italiano",
    "FR": "Francês", "RO": "Romeno", "DE": "Alemão", "NL": "Neerlandês",
    "SV": "Sueco", "NO": "Norueguês", "DA": "Dinamarquês", "IS": "Islandês",
}
FAMILY = {l: ("românica" if l in ("PT", "ES", "IT", "FR", "RO") else "germânica") for l in LANGS}
WARN = "⚠"
# Excel converte true/false em booleano maiúsculo na planilha de origem.
EXCEL_BOOL = {"TRUE": "true", "FALSE": "false"}


def blocks(doc):
    """Percorre parágrafos e tabelas na ordem do documento."""
    for el in doc.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield "p", Paragraph(el, doc)
        elif tag == "tbl":
            yield "t", Table(el, doc)


def rows(table):
    return [[c.text.strip() for c in r.cells] for r in table.rows]


def clean(cell):
    """Remove a marca ⚠ e devolve (texto, tinha_alerta). '—' = vazio no original."""
    flagged = WARN in cell
    text = cell.replace(WARN, "").strip()
    return ("" if text == "—" else text), flagged


def find(files, key):
    hits = [f for f in files if key in f.name]
    if len(hits) != 1:
        sys.exit(f"esperava 1 arquivo contendo {key!r}, achei {len(hits)}")
    return hits[0]


def parse_master(path):
    doc = docx.Document(path)
    words, bands, alerts, ficha, index = {}, [], [], {}, {}
    for kind, b in blocks(doc):
        if kind != "t":
            continue
        r = rows(b)
        head = r[0]
        if head[:3] == ["Nº", "EN", "PT"] and len(head) == 13:
            for row in r[1:]:
                n = int(row[0])
                en = EXCEL_BOOL.get(row[1], row[1])
                tr, flags = {}, []
                for lang, cell in zip(LANGS, row[2:]):
                    tr[lang], f = clean(cell)
                    if f:
                        flags.append(lang)
                words[n] = {"n": n, "en": en, "en_original": row[1], "tr": tr, "flags": flags}
        elif head == ["Faixa", "Nº", "Perfil", "O que predomina"]:
            for row in r[1:]:
                a, z = (int(x) for x in row[1].split("–"))
                bands.append({"faixa": int(row[0]), "de": a, "ate": z, "perfil": row[2], "predomina": row[3]})
        elif head[:3] == ["Nº", "Inglês", "No original"]:
            alerts = [dict(zip(["n", "en", "original", "sugestao", "tipo"], x)) for x in r[1:]]
        elif len(head) == 2 and head[0] == "Fonte":
            ficha = {k: v for k, v in r}
        elif len(head) == 10 and head[1].isdigit():
            for row in r:
                for i in range(0, 10, 2):
                    if row[i]:
                        index[EXCEL_BOOL.get(row[i], row[i])] = int(row[i + 1])
    return words, bands, alerts, ficha, index


def parse_volume(path):
    """Volumes por idioma: tabela de alertas + tabela de duas colunas (8 células)."""
    doc = docx.Document(path)
    alerts, entries, lang = [], {}, None
    for kind, b in blocks(doc):
        if kind != "t":
            continue
        r = rows(b)
        head = r[0]
        if head[:3] == ["Nº", "Inglês", "No original"]:
            alerts = [dict(zip(["n", "en", "original", "sugestao", "tipo"], x)) for x in r[1:]]
        elif len(head) == 8 and head[0] == "Nº":
            lang = head[2]
            for row in r[1:]:
                for half in (row[:4], row[4:]):
                    if half[0]:
                        n = int(half[0])
                        entries[n] = {"en": half[1], "alvo": clean(half[2]), "extra": half[3]}
    return lang, alerts, entries


def parse_ptbr(path):
    """Listas pedagógicas EN→PT-BR (3 colunas) e frases de exemplo 'EN → PT'."""
    doc = docx.Document(path)
    entries, examples, section = {}, [], ""
    for kind, b in blocks(doc):
        if kind == "p":
            t = b.text.strip()
            style = b.style.name if b.style is not None else ""
            if style.startswith("Heading"):
                section = t
            elif t.count("→") == 1 and (not t.startswith("•") or "Exemplo de frase" in t):
                en, pt = (s.strip().strip('"') for s in t.split("→"))
                en = en.lstrip("• ").strip().strip('"')
                if en.lower().startswith("exemplo de frase:"):
                    en = en.split(":", 1)[1].strip().strip('"')
                examples.append({"en": en, "pt": pt, "secao": section})
        else:
            r = rows(b)
            if r[0] == ["#", "Inglês", "Português"]:
                for n, en, pt in r[1:]:
                    entries[int(n)] = {"en": en, "pt": pt}
    return entries, examples


def main(src, out):
    src, out = Path(src), Path(out)
    files = sorted(src.glob("*.docx"))
    out.mkdir(parents=True, exist_ok=True)
    report = {"arquivos": [f.name for f in files], "checagens": [], "divergencias": {}}

    def check(name, ok, detail=None):
        report["checagens"].append({"checagem": name, "ok": bool(ok), **({"detalhe": detail} if detail else {})})

    words, bands, pt_alerts, ficha, index = parse_master(find(files, "Multilingue_2284"))
    check("mestre: 2.284 linhas numeradas 1..2284", sorted(words) == list(range(1, 2285)), len(words))
    total_tr = sum(1 for w in words.values() for v in w["tr"].values() if v)
    check("mestre: equivalentes não vazios", True, f"{total_tr} de {len(words) * 11} células")
    check("índice alfabético cobre todas as palavras",
          all(index.get(w["en"]) == n for n, w in words.items()),
          [w["en"] for n, w in words.items() if index.get(w["en"]) != n][:20] or None)

    def band_of(n):
        return next(b["faixa"] for b in bands if b["de"] <= n <= b["ate"])

    # Volumes por idioma
    alerts_by_lang = {"PT": pt_alerts}
    for key in ("EN-PT", "EN-ES", "EN-FR", "EN-DE"):
        lang, alerts, entries = parse_volume(find(files, key))
        alerts_by_lang[lang] = alerts if lang != "PT" else alerts_by_lang["PT"]
        if lang == "PT":
            check("volume PT: alertas iguais aos do mestre", alerts == pt_alerts)
        diff = []
        for n, e in entries.items():
            w = words[n]
            if EXCEL_BOOL.get(e["en"], e["en"]) != w["en"] or e["alvo"][0] != w["tr"][lang]:
                diff.append({"n": n, "mestre": [w["en"], w["tr"][lang]], "volume": [e["en"], e["alvo"][0]]})
            if lang != "PT" and e["extra"] != w["tr"]["PT"]:
                diff.append({"n": n, "campo": "PT de referência", "mestre": w["tr"]["PT"], "volume": e["extra"]})
        check(f"volume {lang}: 2.284 entradas idênticas ao mestre", len(entries) == 2284 and not diff, diff[:10] or None)
        flagged = {n for n, e in entries.items() if e["alvo"][1]}
        listed = {int(a["n"]) for a in alerts_by_lang[lang]}
        check(f"volume {lang}: marcas ⚠ == lista de alertas", flagged == listed,
              {"so_marca": sorted(flagged - listed), "so_lista": sorted(listed - flagged)} if flagged != listed else None)

    # Listas PT-BR curadas
    ptbr, examples = {}, []
    for key, label in (("Ingles_Pedagogico", "top100"), ("Ingles_Expandido", "top300"), ("Ingles_2000_Completo", "top2000")):
        entries, ex = parse_ptbr(find(files, key))
        for e in ex:
            e["fonte"] = label
        examples += ex
        mismatch = [{"n": n, "lista": e["en"], "mestre": words[n]["en"]}
                    for n, e in entries.items() if EXCEL_BOOL.get(e["en"], e["en"]).lower() != words[n]["en"].lower()]
        check(f"{label}: ordem EN igual ao mestre ({len(entries)} palavras)", not mismatch, mismatch[:10] or None)
        ptbr[label] = entries
    for n in range(1, 101):
        if ptbr["top100"][n]["pt"] != ptbr["top300"][n]["pt"] or ptbr["top300"][n]["pt"] != ptbr["top2000"][n]["pt"]:
            report["divergencias"].setdefault("ptbr_entre_listas", []).append(
                {"n": n, "en": words[n]["en"], **{k: ptbr[k][n]["pt"] for k in ptbr if n in ptbr[k]}})
    for n in range(101, 301):
        if ptbr["top300"][n]["pt"] != ptbr["top2000"][n]["pt"]:
            report["divergencias"].setdefault("ptbr_entre_listas", []).append(
                {"n": n, "en": words[n]["en"], "top300": ptbr["top300"][n]["pt"], "top2000": ptbr["top2000"][n]["pt"]})

    # Mapa final
    alert_idx = {(lang, int(a["n"])): a for lang, al in alerts_by_lang.items() for a in al}
    for lang, al in alerts_by_lang.items():
        bad = [a["n"] for a in al if words[int(a["n"])]["tr"][lang] != a["original"] or words[int(a["n"])]["en"] != EXCEL_BOOL.get(a["en"], a["en"])]
        check(f"alertas {lang}: 'No original' confere com a tabela", not bad, bad or None)

    vocab = []
    for n in sorted(words):
        w = words[n]
        alerts = {lang: {k: a[k] for k in ("original", "sugestao", "tipo")}
                  for lang in alerts_by_lang if (a := alert_idx.get((lang, n)))}
        pt_br = ptbr["top2000"].get(n, {}).get("pt") or ""
        # As listas Top 100/300 trazem mais sentidos para as mesmas palavras.
        ampliado = ptbr["top300"].get(n, {}).get("pt") or ""
        vocab.append({
            "n": n, "faixa": band_of(n), "en": w["en"],
            **({"en_original": w["en_original"]} if w["en_original"] != w["en"] else {}),
            "pt_br": pt_br, **({"pt_br_ampliado": ampliado} if ampliado and ampliado != pt_br else {}),
            "traducoes": w["tr"], "alertas": alerts,
        })

    stats = {
        "palavras": len(vocab), "idiomas": len(LANGS), "equivalentes_nao_vazios": total_tr,
        "celulas_vazias": {l: [v["n"] for v in vocab if not v["traducoes"][l]] for l in LANGS},
        "alertas_por_idioma": {l: len(a) for l, a in alerts_by_lang.items()},
        "tipos_de_alerta_PT": {},
        "palavras_por_faixa": {b["faixa"]: b["ate"] - b["de"] + 1 for b in bands},
        "com_pt_br_curado": sum(1 for v in vocab if v["pt_br"]),
        "frases_exemplo": len(examples),
        "identicas_EN_e_alvo": {l: sum(1 for v in vocab if v["traducoes"][l].lower() == v["en"].lower()) for l in LANGS},
        "pt_br_difere_do_pt_original": sum(1 for v in vocab if v["pt_br"] and v["pt_br"] != v["traducoes"]["PT"]),
    }
    for a in pt_alerts:
        stats["tipos_de_alerta_PT"][a["tipo"]] = stats["tipos_de_alerta_PT"].get(a["tipo"], 0) + 1
    report["estatisticas"] = stats

    meta = {
        "titulo": "Vocabulário Essencial Multilíngue — 2.284 palavras mais frequentes do inglês",
        "autoria": "ODU Creative · Diego Horta · 2026",
        "fonte": ficha.get("Fonte", "Mnemotechnics (mnemotechnics.org)"),
        "ficha_tecnica": ficha,
        "idiomas": {l: {"nome": LANG_NAMES[l], "familia": FAMILY[l]} for l in LANGS},
        "faixas": bands,
        "campos": {
            "n": "posição de frequência (1 = mais usada)",
            "faixa": "faixa de frequência 1–10",
            "en": "palavra-base em inglês (TRUE/FALSE do Excel normalizados para true/false)",
            "pt_br": "tradução PT-BR curada da lista '2000 Palavras' (vazia para 2001–2284)",
            "pt_br_ampliado": "versão com mais sentidos da lista Top 300, quando difere de pt_br",
            "traducoes": "equivalentes do original Mnemotechnics, sem a marca ⚠ ('' = célula vazia '—')",
            "alertas": "por idioma: {original, sugestao, tipo} quando o equivalente é problemático",
        },
    }
    with open(out / "vocabulario.json", "w", encoding="utf-8") as f:
        json.dump({"meta": meta, "palavras": vocab}, f, ensure_ascii=False, indent=1)
    with open(out / "faixas.json", "w", encoding="utf-8") as f:
        json.dump(bands, f, ensure_ascii=False, indent=2)
    with open(out / "validacao.json", "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)

    with open(out / "vocabulario.csv", "w", encoding="utf-8", newline="") as f:
        wr = csv.writer(f)
        wr.writerow(["n", "faixa", "en", "pt_br", "pt_br_ampliado"] + LANGS + ["alertas"])
        for v in vocab:
            wr.writerow([v["n"], v["faixa"], v["en"], v["pt_br"], v.get("pt_br_ampliado", "")] + [v["traducoes"][l] for l in LANGS]
                        + [";".join(sorted(v["alertas"]))])
    with open(out / "alertas.csv", "w", encoding="utf-8", newline="") as f:
        wr = csv.writer(f)
        wr.writerow(["idioma", "n", "en", "original", "sugestao", "tipo"])
        for lang, al in alerts_by_lang.items():
            for a in al:
                wr.writerow([lang, a["n"], EXCEL_BOOL.get(a["en"], a["en"]), a["original"], a["sugestao"], a["tipo"]])
    with open(out / "frases_exemplo.csv", "w", encoding="utf-8", newline="") as f:
        wr = csv.writer(f)
        wr.writerow(["fonte", "secao", "en", "pt"])
        for e in examples:
            wr.writerow([e["fonte"], e["secao"], e["en"], e["pt"]])

    # Baralhos Anki (TSV): frente EN, verso = forma recomendada no idioma-alvo
    anki = out.parent / "anki"
    anki.mkdir(exist_ok=True)
    for lang in LANGS:
        with open(anki / f"EN-{lang}.tsv", "w", encoding="utf-8", newline="") as f:
            f.write("#separator:tab\n#html:true\n#columns:Frente\tVerso\tTags\n")
            for v in vocab:
                tr = v["traducoes"][lang]
                a = v["alertas"].get(lang)
                if lang == "PT" and v["pt_br"]:
                    back = v.get("pt_br_ampliado") or v["pt_br"]
                elif a:
                    back = a["sugestao"]
                else:
                    back = tr
                if not back:
                    continue
                if lang != "PT" and v["pt_br"]:
                    back += f"<br><small>PT: {v['pt_br']}</small>"
                tags = f"faixa{v['faixa']:02d} n{v['n']:04d}" + (" alerta" if a else "")
                f.write(f"{v['en']}\t{back}\t{tags}\n")

    ok = sum(c["ok"] for c in report["checagens"])
    print(f"{ok}/{len(report['checagens'])} checagens ok")
    for c in report["checagens"]:
        print(("OK  " if c["ok"] else "FALHA ") + c["checagem"], "" if c["ok"] else json.dumps(c.get("detalhe"), ensure_ascii=False)[:600])
    print(json.dumps(stats | {"celulas_vazias": {k: len(v) for k, v in stats["celulas_vazias"].items()}}, ensure_ascii=False))


if __name__ == "__main__":
    main(*sys.argv[1:3])
